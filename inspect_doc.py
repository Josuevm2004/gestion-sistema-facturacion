from docx import Document
p='Contexto_y_mapeo_BPMN_Miquipu.docx'; d=Document(p)
for i,x in enumerate(d.paragraphs): print(i,repr(x.text))
for ti,t in enumerate(d.tables): print('TABLE',ti);[print([c.text for c in r.cells]) for r in t.rows]
